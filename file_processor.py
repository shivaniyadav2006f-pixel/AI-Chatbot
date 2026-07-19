try:
    # prefer explicit PdfReader import
    from PyPDF2 import PdfReader  # type: ignore[import-not-found]
except ImportError:
    try:
        # some environments provide pypdf package
        from pypdf import PdfReader  # type: ignore[import-not-found]
    except ImportError:
        PdfReader = None
from docx import Document
from pptx import Presentation


# -------------------------
# PDF Reader
# -------------------------
def read_pdf(file):

    text = ""

    if PdfReader is None:
        raise ImportError("PyPDF2 or pypdf is required to read PDF files")

    reader = PdfReader(file)

    for page in reader.pages:
        text += page.extract_text() or ""

    return text



# -------------------------
# DOCX Reader
# -------------------------
def read_docx(file):

    doc = Document(file)

    text = ""

    for para in doc.paragraphs:
        text += para.text + "\n"

    return text



# -------------------------
# PPT Reader
# -------------------------
def read_ppt(file):

    prs = Presentation(file)

    text = ""

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text += shape.text + "\n"

    return text



# -------------------------
# TXT Reader
# -------------------------
def read_txt(file):
    data = file.read()
    if isinstance(data, bytes):
        return data.decode("utf-8")
    return data



# -------------------------
# Main Processor
# -------------------------
def process_file(file):

    filename = file.name.lower()


    if filename.endswith(".pdf"):

        return read_pdf(file)


    elif filename.endswith(".docx"):

        return read_docx(file)


    elif filename.endswith(".pptx"):

        return read_ppt(file)


    elif filename.endswith(".txt"):

        return read_txt(file)


    else:

        return "Unsupported file"